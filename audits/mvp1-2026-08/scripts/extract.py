import os, glob, json, warnings
warnings.filterwarnings('ignore')
SP=os.path.dirname(os.path.abspath(__file__))
R=os.path.join(SP,'reports'); T=os.path.join(SP,'text')

def pdf(p):
    import pypdf
    out=[]
    try:
        r=pypdf.PdfReader(p)
        for i,pg in enumerate(r.pages):
            out.append(f"--- page {i+1} ---\n"+(pg.extract_text() or ''))
    except Exception as e: out.append(f"[PDF ERROR {e}]")
    return "\n".join(out)

def xlsx(p):
    import openpyxl
    out=[]
    try:
        wb=openpyxl.load_workbook(p, data_only=True)
        for ws in wb.worksheets:
            out.append(f"=== SHEET: {ws.title} ({ws.max_row}x{ws.max_column}) ===")
            for row in ws.iter_rows(values_only=True):
                cells=[str(c) for c in row if c is not None and str(c).strip()]
                if cells: out.append(" | ".join(cells))
    except Exception as e: out.append(f"[XLSX ERROR {e}]")
    return "\n".join(out)

def xls(p):
    import xlrd
    out=[]
    try:
        wb=xlrd.open_workbook(p)
        for ws in wb.sheets():
            out.append(f"=== SHEET: {ws.name} ({ws.nrows}x{ws.ncols}) ===")
            for i in range(ws.nrows):
                cells=[str(c.value) for c in ws.row(i) if str(c.value).strip()]
                if cells: out.append(" | ".join(cells))
    except Exception as e: out.append(f"[XLS ERROR {e}]")
    return "\n".join(out)

def docx_(p):
    import docx
    out=[]
    try:
        d=docx.Document(p)
        for par in d.paragraphs:
            if par.text.strip(): out.append(par.text)
        for ti,tb in enumerate(d.tables):
            out.append(f"=== TABLE {ti+1} ===")
            for row in tb.rows:
                out.append(" | ".join(c.text.strip() for c in row.cells))
    except Exception as e: out.append(f"[DOCX ERROR {e}]")
    return "\n".join(out)

for f in sorted(glob.glob(os.path.join(R,'*'))):
    b=os.path.basename(f); ext=b.rsplit('.',1)[-1].lower()
    fn={'pdf':pdf,'xlsx':xlsx,'xls':xls,'docx':docx_}.get(ext)
    if not fn: continue
    txt=fn(f)
    open(os.path.join(T,b+'.txt'),'w').write(txt)
    print(f"{b:55} {len(txt):>8} chars")
